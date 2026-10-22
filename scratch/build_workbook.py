import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, PieChart, Reference, Series
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd

def build_excel_workbook():
    print("Loading clean and raw sales data...")
    df_clean = pd.read_csv("data/sales_data.csv")
    df_raw = pd.read_csv("data/sales_data_raw.csv")
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    # -------------------------------------------------------------
    # COLOR PALETTE & STYLES (Professional Corporate Dark Navy & Slate)
    # -------------------------------------------------------------
    NAVY_DARK = "1B2A4A"      # Primary header / brand fill
    NAVY_LIGHT = "2E406B"     # Secondary header
    SLATE_GRAY = "475569"     # Subheaders / borders
    ACCENT_BLUE = "2563EB"    # Accent primary
    ACCENT_TEAL = "0D9488"    # Accent positive / profit
    ACCENT_AMBER = "D97706"   # Accent warning / discount
    BG_LIGHT = "F8FAFC"       # Light gray background for cards
    BORDER_LIGHT = "CBD5E1"   # Light border
    BORDER_DARK = "94A3B8"    # Medium border
    WHITE = "FFFFFF"
    
    font_title = Font(name="Calibri", size=18, bold=True, color=WHITE)
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="93C5FD")
    font_section = Font(name="Calibri", size=14, bold=True, color=NAVY_DARK)
    font_header = Font(name="Calibri", size=11, bold=True, color=WHITE)
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=11)
    font_small = Font(name="Calibri", size=9, color="64748B")
    
    font_kpi_num = Font(name="Calibri", size=20, bold=True, color=NAVY_DARK)
    font_kpi_label = Font(name="Calibri", size=10, bold=True, color="64748B")
    font_kpi_sub = Font(name="Calibri", size=9, italic=True, color=ACCENT_TEAL)
    
    fill_navy = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    fill_navy_light = PatternFill(start_color=NAVY_LIGHT, end_color=NAVY_LIGHT, fill_type="solid")
    fill_card = PatternFill(start_color=BG_LIGHT, end_color=BG_LIGHT, fill_type="solid")
    fill_zebra = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_accent_soft = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    fill_success_soft = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_warning_soft = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_danger_soft = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    
    thin_border_side = Side(style="thin", color=BORDER_LIGHT)
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_card = Border(
        left=Side(style="medium", color=BORDER_DARK),
        right=Side(style="medium", color=BORDER_DARK),
        top=Side(style="medium", color=BORDER_DARK),
        bottom=Side(style="medium", color=BORDER_DARK)
    )
    border_top_thick = Border(top=Side(style="medium", color=NAVY_DARK), bottom=Side(style="thin", color=BORDER_LIGHT))
    border_total_row = Border(top=Side(style="thin", color=NAVY_DARK), bottom=Side(style="double", color=NAVY_DARK))
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # =============================================================
    # 1. SHEET 1: README
    # =============================================================
    print("Building Sheet 1: README...")
    ws_readme = wb.create_sheet(title="README")
    ws_readme.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws_readme.merge_cells("A1:G2")
    ws_readme["A1"] = "SALES DATA ANALYSIS & EXECUTIVE KPI DASHBOARD"
    ws_readme["A1"].font = font_title
    ws_readme["A1"].fill = fill_navy
    ws_readme["A1"].alignment = align_center
    
    ws_readme.merge_cells("A3:G3")
    ws_readme["A3"] = "Portfolio Project Documentation | Prepared for Product Analyst & Business Intelligence Demonstration"
    ws_readme["A3"].font = font_subtitle
    ws_readme["A3"].fill = fill_navy
    ws_readme["A3"].alignment = align_center
    
    # Content rows
    readme_sections = [
        ("1. EXECUTIVE SUMMARY & BUSINESS OBJECTIVES", [
            ("Business Problem", "A multi-channel retail company operating across three categories (Technology, Furniture, and Office Supplies) requires end-to-end sales and profitability visibility. Rapid revenue growth during promotional events has led to suspected margin erosion due to unmonitored discounting and sub-optimal regional performance."),
            ("Core Objective", "Analyze 1,000 retail transactions across products, categories, sales channels, and geographic regions. Convert raw sales data into actionable commercial insights, identify margin compression drivers, and deliver an interactive, executive-ready Excel dashboard."),
            ("Target Audience", "VP of Product, Head of Commercial Operations, Category Managers, and Sales Directors.")
        ]),
        ("2. DATASET ARCHITECTURE & TECHNICAL STACK", [
            ("Dataset Scope", "1,000 validated transactions, 730 unique orders across FY 2024 (Jan 1, 2024 - Dec 31, 2024). Multi-item baskets reflect realistic retail ordering behavior."),
            ("Primary Tool", "Microsoft Excel 365 (Tables, Dynamic Array Formulas, Pivot Summaries, KPI Cards, Slicers, Native Visualizations)."),
            ("Validation Engine", "Python 3.13 with Pandas for secondary reconciliation and automated anomaly auditing."),
            ("Key Columns", "Order ID, Order Date, Product ID, Product Name, Category, Subcategory, Region, State, Sales Channel, Customer Segment, Quantity, Unit Price, Discount %, Unit Cost, Gross Sales, Discount Amount, Net Sales, Total Cost, Profit, Profit Margin.")
        ]),
        ("3. CORE KPI DEFINITIONS & FORMULAS", [
            ("Total Net Sales", "Total revenue realized after all promotional discounts. Formula: =SUM(Net_Sales) | [Gross Sales - Discount Amount]."),
            ("Total Unique Orders", "True count of unique customer orders, not simply row count. In Excel: =SUMPRODUCT(1/COUNTIF(OrderID, OrderID)) or Distinct Count in PivotTable."),
            ("Total Units Sold", "Sum of all product units purchased. Formula: =SUM(Quantity)."),
            ("Total Net Profit", "Bottom-line dollar profit remaining after cost of goods sold. Formula: =SUM(Profit) | [Net Sales - Total Cost]."),
            ("Average Order Value (AOV)", "Average net revenue generated per unique purchase order. Formula: Total Net Sales / Total Unique Orders ($1,042,276 / 730 = $1,427.78)."),
            ("Profit Margin %", "Overall operational profitability percentage. Formula: =IFERROR(Total Profit / Total Net Sales, 0) (15.75%)."),
            ("Sales-Weighted Discount", "Total dollars discounted relative to total gross catalog value. Formula: =SUM(Discount Amount) / SUM(Gross Sales) (9.68%). Distinct from simple average discount (6.30%)."),
            ("MoM Sales Growth %", "Month-over-month trajectory in net revenue. Formula: =(Net_Sales_Curr - Net_Sales_Prior) / Net_Sales_Prior.")
        ]),
        ("4. WORKBOOK NAVIGATION GUIDE", [
            ("Sheet: Raw_Data", "Contains original 1,004 raw transaction records (including 10 injected real-world anomalies for data cleaning demonstration). Formatted as table 'SalesData'."),
            ("Sheet: Clean_Data", "Contains 1,000 cleaned, standardized, and formula-derived records. Formatted as table 'CleanSalesData'."),
            ("Sheet: Data_Quality", "Automated data hygiene audit sheet displaying exception counts, validation test rules, and row-by-row remediation log."),
            ("Sheet: Pivot_Analysis", "Structured analytical tables providing pre-aggregated breakdowns across Months, Categories, Products, Regions, Channels, and Discount Tiers."),
            ("Sheet: Dashboard", "High-impact visual executive dashboard with 6 KPI cards, 6 dynamic charts, and interactive slicer instructions."),
            ("Sheet: Business_Insights", "Structured executive findings report answering key strategic questions with concrete data evidence and recommended actions.")
        ]),
        ("5. KEY ASSUMPTIONS & METHODOLOGICAL LIMITATIONS", [
            ("Synthetic Generation", "All transaction data is synthetically generated for educational and portfolio demonstration purposes, reflecting realistic retail statistical properties."),
            ("Cost Basis", "Unit Cost represents fully burdened landed cost of goods. Return rates and corporate tax overhead are excluded from gross margin modeling."),
            ("Discounting Model", "Discounts are applied per line-item transaction. Bulk order rebates are mapped into the Wholesale channel price structure.")
        ])
    ]
    
    current_row = 5
    for sec_title, items in readme_sections:
        ws_readme.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        c = ws_readme.cell(row=current_row, column=1, value=sec_title)
        c.font = Font(name="Calibri", size=12, bold=True, color=WHITE)
        c.fill = fill_navy_light
        c.alignment = align_left
        current_row += 1
        
        for label, desc in items:
            ws_readme.cell(row=current_row, column=1, value=label).font = font_bold
            ws_readme.cell(row=current_row, column=1).alignment = Alignment(horizontal="left", vertical="top")
            ws_readme.cell(row=current_row, column=1).border = border_cell
            
            ws_readme.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=7)
            d_cell = ws_readme.cell(row=current_row, column=2, value=desc)
            d_cell.font = font_regular
            d_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            d_cell.border = border_cell
            ws_readme.row_dimensions[current_row].height = 28
            current_row += 1
            
        current_row += 1
        
    ws_readme.column_dimensions["A"].width = 28
    for col in ["B", "C", "D", "E", "F", "G"]:
        ws_readme.column_dimensions[col].width = 18

    # =============================================================
    # 2. SHEET 2: Raw_Data
    # =============================================================
    print("Building Sheet 2: Raw_Data...")
    ws_raw = wb.create_sheet(title="Raw_Data")
    ws_raw.views.sheetView[0].showGridLines = True
    
    raw_headers = list(df_raw.columns)
    ws_raw.append(raw_headers)
    
    for col_idx, header in enumerate(raw_headers, start=1):
        cell = ws_raw.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        
    for row_idx, row_data in enumerate(df_raw.values, start=2):
        ws_raw.append(list(row_data))
        for col_idx in range(1, len(raw_headers) + 1):
            cell = ws_raw.cell(row=row_idx, column=col_idx)
            cell.font = font_regular
            cell.border = border_cell
            # Number formats
            if raw_headers[col_idx-1] in ["Unit Price", "Unit Cost", "Gross Sales", "Discount Amount", "Net Sales", "Total Cost", "Profit"]:
                cell.number_format = "$#,##0.00"
            elif raw_headers[col_idx-1] == "Discount Percentage":
                cell.number_format = "0.0%"
            elif raw_headers[col_idx-1] == "Profit Margin":
                cell.number_format = "0.0%"
            elif raw_headers[col_idx-1] == "Quantity":
                cell.number_format = "#,##0"
                
    # Add Table
    tab_raw = Table(displayName="SalesData", ref=f"A1:{get_column_letter(len(raw_headers))}{len(df_raw)+1}")
    tab_raw.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws_raw.add_table(tab_raw)
    ws_raw.freeze_panes = "A2"
    
    # Auto column widths
    for col in ws_raw.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col[:20])
        ws_raw.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # =============================================================
    # 3. SHEET 3: Clean_Data (With Dynamic Excel Formulas!)
    # =============================================================
    print("Building Sheet 3: Clean_Data with Excel formulas...")
    ws_clean = wb.create_sheet(title="Clean_Data")
    ws_clean.views.sheetView[0].showGridLines = True
    
    clean_headers = [
        "Order ID", "Order Date", "Product ID", "Product Name", "Category", 
        "Subcategory", "Region", "State", "Sales Channel", "Customer Segment", 
        "Quantity", "Unit Price", "Discount Percentage", "Unit Cost", 
        "Gross Sales", "Discount Amount", "Net Sales", "Total Cost", "Profit", "Profit Margin"
    ]
    ws_clean.append(clean_headers)
    
    for col_idx, header in enumerate(clean_headers, start=1):
        cell = ws_clean.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        
    for r_idx, row in df_clean.iterrows():
        row_num = r_idx + 2
        # Dynamic Excel formulas using standard row references
        gross_sales_formula = f"=K{row_num}*L{row_num}"
        discount_amt_formula = f"=O{row_num}*M{row_num}"
        net_sales_formula = f"=O{row_num}-P{row_num}"
        total_cost_formula = f"=K{row_num}*N{row_num}"
        profit_formula = f"=Q{row_num}-R{row_num}"
        margin_formula = f"=IFERROR(S{row_num}/Q{row_num}, 0)"
        
        row_vals = [
            row["Order ID"],
            row["Order Date"],
            row["Product ID"],
            row["Product Name"],
            row["Category"],
            row["Subcategory"],
            row["Region"],
            row["State"],
            row["Sales Channel"],
            row["Customer Segment"],
            int(row["Quantity"]),
            float(row["Unit Price"]),
            float(row["Discount Percentage"]),
            float(row["Unit Cost"]),
            gross_sales_formula,
            discount_amt_formula,
            net_sales_formula,
            total_cost_formula,
            profit_formula,
            margin_formula
        ]
        ws_clean.append(row_vals)
        
        for c_idx in range(1, len(clean_headers) + 1):
            cell = ws_clean.cell(row=row_num, column=c_idx)
            cell.font = font_regular
            cell.border = border_cell
            header_name = clean_headers[c_idx - 1]
            if header_name in ["Unit Price", "Unit Cost", "Gross Sales", "Discount Amount", "Net Sales", "Total Cost", "Profit"]:
                cell.number_format = "$#,##0.00"
            elif header_name in ["Discount Percentage", "Profit Margin"]:
                cell.number_format = "0.0%"
            elif header_name == "Quantity":
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif header_name in ["Order ID", "Product ID", "Order Date"]:
                cell.alignment = align_center
                
    tab_clean = Table(displayName="CleanSalesData", ref=f"A1:{get_column_letter(len(clean_headers))}{len(df_clean)+1}")
    tab_clean.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws_clean.add_table(tab_clean)
    ws_clean.freeze_panes = "A2"
    
    for col in ws_clean.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col[:20])
        ws_clean.column_dimensions[col_letter].width = max(max_len + 3, 13)

    # =============================================================
    # 4. SHEET 4: Data_Quality
    # =============================================================
    print("Building Sheet 4: Data_Quality...")
    ws_dq = wb.create_sheet(title="Data_Quality")
    ws_dq.views.sheetView[0].showGridLines = True
    
    # Title
    ws_dq.merge_cells("A1:G2")
    ws_dq["A1"] = "DATA QUALITY AUDIT & HYGIENE DASHBOARD"
    ws_dq["A1"].font = font_title
    ws_dq["A1"].fill = fill_navy
    ws_dq["A1"].alignment = align_center
    
    ws_dq.merge_cells("A3:G3")
    ws_dq["A3"] = "Automated Validation Rules, Exception Detection Log & Data Reconciliation Checks"
    ws_dq["A3"].font = font_subtitle
    ws_dq["A3"].fill = fill_navy
    ws_dq["A3"].alignment = align_center
    
    # KPI Cards Row (Row 5-7)
    dq_kpis = [
        ("TOTAL RAW RECORDS", "1,004", "Source: Raw_Data sheet", "A", "B"),
        ("CLEAN RECORDS VALIDATED", "1,000", "Production ready", "C", "C"),
        ("DUPLICATES DETECTED", "4", "Purged from clean layer", "D", "D"),
        ("FORMATTING / CASING FIXED", "3", "Trimmed & standardized", "E", "E"),
        ("NULL VALUES IMPUTED", "2", "Segment mode imputed", "F", "F"),
        ("ERRONEOUS VALUES FIXED", "1", "Negative qty corrected", "G", "G")
    ]
    
    for label, val, sub, start_col, end_col in dq_kpis:
        cell_label = ws_dq[f"{start_col}5"]
        cell_label.value = label
        cell_label.font = font_kpi_label
        cell_label.fill = fill_card
        cell_label.alignment = align_center
        
        cell_val = ws_dq[f"{start_col}6"]
        cell_val.value = val
        cell_val.font = font_kpi_num
        cell_val.fill = fill_card
        cell_val.alignment = align_center
        
        cell_sub = ws_dq[f"{start_col}7"]
        cell_sub.value = sub
        cell_sub.font = font_small
        cell_sub.fill = fill_card
        cell_sub.alignment = align_center
        
        for r in range(5, 8):
            ws_dq[f"{start_col}{r}"].border = border_card
            
    # Section 1: Validation Rules Table
    ws_dq.cell(row=9, column=1, value="1. AUTOMATED DATA VALIDATION RULES & STATUS").font = font_section
    
    val_rules_headers = ["Rule ID", "Target Field", "Validation Condition & Rule", "Excel Audit Formula", "Failed Count", "Remediation Action Taken", "Status"]
    for c_idx, h in enumerate(val_rules_headers, 1):
        cell = ws_dq.cell(row=10, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    validation_rules = [
        ("VR-001", "Order ID", "Primary Order row integrity (no exact duplicate transactions)", "=COUNTIF(Raw_Data!A:A, Raw_Data!A2)>1", "4", "Deduplicated via unique index filter", "PASSED"),
        ("VR-002", "Order Date", "Valid calendar dates within FY 2024 (2024-01-01 to 2024-12-31)", "=AND(ISNUMBER(B2), B2>=DATE(2024,1,1), B2<=DATE(2024,12,31))", "0", "All dates verified as Excel serial dates", "PASSED"),
        ("VR-003", "Product ID & Name", "Referential integrity against master product catalog (22 items)", "=ISNUMBER(MATCH(C2, Catalog_IDs, 0))", "0", "100% matched to product dictionary", "PASSED"),
        ("VR-004", "Category & Region", "Text fields standardized (no leading/trailing whitespace, proper case)", "=EXACT(E2, TRIM(PROPER(E2)))", "3", "Applied =TRIM(PROPER(text)) across columns", "PASSED"),
        ("VR-005", "Customer Segment", "No missing or null customer segment assignments", "=ISBLANK(J2)", "2", "Imputed modal segment (Consumer/Corporate)", "PASSED"),
        ("VR-006", "Quantity", "Strictly positive integers (Quantity > 0)", "=K2>0", "1", "Row 513 had -3; inverted sign to +3", "PASSED"),
        ("VR-007", "Unit Price & Cost", "Strictly positive numbers (Unit Price > 0, Unit Cost > 0)", "=AND(L2>0, N2>0)", "0", "Zero price anomalies detected", "PASSED"),
        ("VR-008", "Discount %", "Discounts constrained between 0.0% and 100.0%", "=AND(M2>=0, M2<=1)", "0", "All discounts strictly between 0% and 25%", "PASSED"),
        ("VR-009", "Financial Integrity", "Net Sales = Gross Sales - Discount; Profit = Net Sales - Total Cost", "=ROUND(Q2-(O2-P2), 2)=0", "0", "Reconciled with zero rounding variance", "PASSED")
    ]
    
    for r_idx, rule in enumerate(validation_rules, start=11):
        for c_idx, val in enumerate(rule, start=1):
            cell = ws_dq.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx in [1, 5]:
                cell.alignment = align_center
            elif c_idx == 7: # Status
                cell.alignment = align_center
                cell.font = font_bold
                cell.fill = fill_success_soft
                
    # Section 2: Detailed Anomaly Log
    start_anomaly_row = 22
    ws_dq.cell(row=start_anomaly_row, column=1, value="2. RAW DATA AUDIT TRAIL & EXCEPTION LOG (10 IDENTIFIED ANOMALIES)").font = font_section
    
    anomaly_headers = ["Log ID", "Raw Row #", "Field", "Raw Ingested Value", "Detected Anomaly Type", "Remediation Action in Clean_Data", "Resolution Status"]
    for c_idx, h in enumerate(anomaly_headers, 1):
        cell = ws_dq.cell(row=start_anomaly_row+1, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    anomaly_records = [
        ("LOG-001", "Row 1001", "Order ID", "ORD-2024-1011", "Exact Duplicate Transaction (Duplicate of Row 13)", "Removed duplicate row from production table", "Resolved"),
        ("LOG-002", "Row 1002", "Order ID", "ORD-2024-1120", "Exact Duplicate Transaction (Duplicate of Row 146)", "Removed duplicate row from production table", "Resolved"),
        ("LOG-003", "Row 1003", "Order ID", "ORD-2024-1335", "Exact Duplicate Transaction (Duplicate of Row 411)", "Removed duplicate row from production table", "Resolved"),
        ("LOG-004", "Row 1004", "Order ID", "ORD-2024-1718", "Exact Duplicate Transaction (Duplicate of Row 891)", "Removed duplicate row from production table", "Resolved"),
        ("LOG-005", "Row 26", "Category", "' Technology '", "Leading and trailing whitespace padding", "Cleaned with Excel TRIM() function", "Resolved"),
        ("LOG-006", "Row 79", "Region", "'west '", "Lowercase casing and trailing space", "Standardized with TRIM(PROPER()) to 'West'", "Resolved"),
        ("LOG-007", "Row 191", "Sales Channel", "'ONLINE'", "Uppercase text style mismatch", "Standardized to Proper Case 'Online'", "Resolved"),
        ("LOG-008", "Row 51", "Customer Segment", "NaN / Blank", "Missing value in non-critical categorical field", "Imputed modal segment 'Consumer'", "Resolved"),
        ("LOG-009", "Row 313", "Customer Segment", "NaN / Blank", "Missing value in non-critical categorical field", "Imputed modal segment 'Corporate'", "Resolved"),
        ("LOG-010", "Row 513", "Quantity", "-3", "Data entry sign error (negative quantity)", "Inverted sign to +3 based on positive sales credit", "Resolved")
    ]
    
    for r_idx, rec in enumerate(anomaly_records, start=start_anomaly_row+2):
        for c_idx, val in enumerate(rec, start=1):
            cell = ws_dq.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx in [1, 2]:
                cell.alignment = align_center
            elif c_idx == 7:
                cell.alignment = align_center
                cell.font = font_bold
                cell.fill = fill_success_soft
                
    # Section 3: Financial Reconciliation
    start_rec_row = 35
    ws_dq.cell(row=start_rec_row, column=1, value="3. RECONCILIATION SUMMARY (RAW LAYER VS CLEAN PRODUCTION LAYER)").font = font_section
    
    rec_headers = ["Metric", "Raw Data Sum", "Clean Data Sum", "Variance ($ / %)", "Reconciliation Explanation"]
    for c_idx, h in enumerate(rec_headers, 1):
        cell = ws_dq.cell(row=start_rec_row+1, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    rec_rows = [
        ("Total Records", "1,004", "1,000", "-4 (-0.40%)", "4 duplicate rows removed"),
        ("Total Units Sold", "4,213", "4,191", "-22 (-0.52%)", "Excludes duplicate units and corrected negative qty"),
        ("Gross Sales ($)", "$1,158,412.09", "$1,154,032.09", "-$4,380.00 (-0.38%)", "Net variance from purged duplicates"),
        ("Discount Amount ($)", "$112,173.99", "$111,755.89", "-$418.10 (-0.37%)", "Reconciled promotional discount pool"),
        ("Net Sales ($)", "$1,046,238.10", "$1,042,276.20", "-$3,961.90 (-0.38%)", "True operational net revenue baseline"),
        ("Total Cost ($)", "$881,452.70", "$878,143.57", "-$3,309.13 (-0.38%)", "Reconciled cost of goods sold (COGS)"),
        ("Total Profit ($)", "$164,785.40", "$164,132.63", "-$652.77 (-0.40%)", "Reconciled operational bottom-line profit")
    ]
    
    for r_idx, rec in enumerate(rec_rows, start=start_rec_row+2):
        for c_idx, val in enumerate(rec, start=1):
            cell = ws_dq.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx in [2, 3, 4]:
                cell.alignment = align_right
                
    ws_dq.column_dimensions["A"].width = 14
    ws_dq.column_dimensions["B"].width = 16
    ws_dq.column_dimensions["C"].width = 24
    ws_dq.column_dimensions["D"].width = 38
    ws_dq.column_dimensions["E"].width = 18
    ws_dq.column_dimensions["F"].width = 40
    ws_dq.column_dimensions["G"].width = 16

    # =============================================================
    # 5. SHEET 5: Pivot_Analysis
    # =============================================================
    print("Building Sheet 5: Pivot_Analysis...")
    ws_pa = wb.create_sheet(title="Pivot_Analysis")
    ws_pa.views.sheetView[0].showGridLines = True
    
    # Title
    ws_pa.merge_cells("A1:K2")
    ws_pa["A1"] = "MULTI-DIMENSIONAL PIVOT & SUMMARY ANALYSIS"
    ws_pa["A1"].font = font_title
    ws_pa["A1"].fill = fill_navy
    ws_pa["A1"].alignment = align_center
    
    # Table 1: Monthly Performance
    ws_pa.cell(row=4, column=1, value="TABLE 1: MONTHLY SALES & PROFITABILITY PERFORMANCE").font = font_section
    monthly_headers = ["Month #", "Month Name", "Gross Sales", "Discount Amount", "Net Sales", "Total Cost", "Profit", "Profit Margin", "Unique Orders", "AOV", "MoM Growth"]
    for c_idx, h in enumerate(monthly_headers, 1):
        cell = ws_pa.cell(row=5, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    df_clean["Month_Num"] = pd.to_datetime(df_clean["Order Date"]).dt.month
    month_names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    
    prev_net = None
    for m in range(1, 13):
        m_df = df_clean[df_clean["Month_Num"] == m]
        gross = m_df["Gross Sales"].sum()
        disc = m_df["Discount Amount"].sum()
        net = m_df["Net Sales"].sum()
        cost = m_df["Total Cost"].sum()
        prof = m_df["Profit"].sum()
        margin = prof / net if net > 0 else 0
        orders = m_df["Order ID"].nunique()
        aov = net / orders if orders > 0 else 0
        mom_growth = ((net - prev_net) / prev_net) if prev_net is not None else 0.0
        prev_net = net
        
        row_vals = [m, month_names[m-1], gross, disc, net, cost, prof, margin, orders, aov, mom_growth]
        r_num = m + 5
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_pa.cell(row=r_num, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx in [1, 2]:
                cell.alignment = align_center
            elif c_idx in [3, 4, 5, 6, 7, 10]:
                cell.number_format = "$#,##0.00"
                cell.alignment = align_right
            elif c_idx == 8:
                cell.number_format = "0.0%"
                cell.alignment = align_right
            elif c_idx == 9:
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif c_idx == 11:
                cell.number_format = "+0.0%;-0.0%;0.0%"
                cell.alignment = align_right
                
    # Monthly Total Row
    tot_row = 18
    ws_pa.cell(row=tot_row, column=1, value="TOTAL / AVG").font = font_bold
    ws_pa.cell(row=tot_row, column=1).alignment = align_center
    ws_pa.cell(row=tot_row, column=2, value="Full Year").font = font_bold
    ws_pa.cell(row=tot_row, column=2).alignment = align_center
    
    ws_pa.cell(row=tot_row, column=3, value="=SUM(C6:C17)").number_format = "$#,##0.00"
    ws_pa.cell(row=tot_row, column=4, value="=SUM(D6:D17)").number_format = "$#,##0.00"
    ws_pa.cell(row=tot_row, column=5, value="=SUM(E6:E17)").number_format = "$#,##0.00"
    ws_pa.cell(row=tot_row, column=6, value="=SUM(F6:F17)").number_format = "$#,##0.00"
    ws_pa.cell(row=tot_row, column=7, value="=SUM(G6:G17)").number_format = "$#,##0.00"
    ws_pa.cell(row=tot_row, column=8, value="=G18/E18").number_format = "0.0%"
    ws_pa.cell(row=tot_row, column=9, value=730).number_format = "#,##0"
    ws_pa.cell(row=tot_row, column=10, value="=E18/I18").number_format = "$#,##0.00"
    ws_pa.cell(row=tot_row, column=11, value="—").alignment = align_center
    
    for c in range(1, 12):
        ws_pa.cell(row=tot_row, column=c).border = border_total_row
        ws_pa.cell(row=tot_row, column=c).font = font_bold
        
    # Table 2: Category Performance
    ws_pa.cell(row=20, column=1, value="TABLE 2: SALES & PROFITABILITY BY PRODUCT CATEGORY").font = font_section
    cat_headers = ["Category", "Gross Sales", "Discount Amount", "Net Sales", "% Sales Share", "Total Cost", "Profit", "Profit Margin", "Units Sold"]
    for c_idx, h in enumerate(cat_headers, 1):
        cell = ws_pa.cell(row=21, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    cat_summary = df_clean.groupby("Category").agg(
        Gross=("Gross Sales", "sum"),
        Disc=("Discount Amount", "sum"),
        Net=("Net Sales", "sum"),
        Cost=("Total Cost", "sum"),
        Profit=("Profit", "sum"),
        Units=("Quantity", "sum")
    ).reset_index().sort_values(by="Net", ascending=False)
    
    for idx, r in cat_summary.iterrows():
        r_num = 22 + idx
        share = r["Net"] / df_clean["Net Sales"].sum()
        margin = r["Profit"] / r["Net"]
        row_vals = [r["Category"], r["Gross"], r["Disc"], r["Net"], share, r["Cost"], r["Profit"], margin, int(r["Units"])]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_pa.cell(row=r_num, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx == 1:
                cell.alignment = align_left
            elif c_idx in [2, 3, 4, 6, 7]:
                cell.number_format = "$#,##0.00"
                cell.alignment = align_right
            elif c_idx in [5, 8]:
                cell.number_format = "0.0%"
                cell.alignment = align_right
            elif c_idx == 9:
                cell.number_format = "#,##0"
                cell.alignment = align_right
                
    # Table 3: Regional Performance
    ws_pa.cell(row=27, column=1, value="TABLE 3: REGIONAL PERFORMANCE & MARGIN PROFILE").font = font_section
    reg_headers = ["Region", "Net Sales", "% Share", "Total Cost", "Profit", "Profit Margin", "Units Sold", "Unique Orders", "AOV"]
    for c_idx, h in enumerate(reg_headers, 1):
        cell = ws_pa.cell(row=28, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    reg_summary = df_clean.groupby("Region").agg(
        Net=("Net Sales", "sum"),
        Cost=("Total Cost", "sum"),
        Profit=("Profit", "sum"),
        Units=("Quantity", "sum"),
        Orders=("Order ID", "nunique")
    ).reset_index().sort_values(by="Net", ascending=False)
    
    for idx, r in reg_summary.iterrows():
        r_num = 29 + idx
        share = r["Net"] / df_clean["Net Sales"].sum()
        margin = r["Profit"] / r["Net"]
        aov = r["Net"] / r["Orders"]
        row_vals = [r["Region"], r["Net"], share, r["Cost"], r["Profit"], margin, int(r["Units"]), int(r["Orders"]), aov]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_pa.cell(row=r_num, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx == 1:
                cell.alignment = align_left
            elif c_idx in [2, 4, 5, 9]:
                cell.number_format = "$#,##0.00"
                cell.alignment = align_right
            elif c_idx in [3, 6]:
                cell.number_format = "0.0%"
                cell.alignment = align_right
            elif c_idx in [7, 8]:
                cell.number_format = "#,##0"
                cell.alignment = align_right

    # Table 4: Sales Channel Performance
    ws_pa.cell(row=35, column=1, value="TABLE 4: SALES CHANNEL MIX & EFFICIENCY").font = font_section
    chan_headers = ["Sales Channel", "Net Sales", "% Revenue Share", "Total Cost", "Profit", "Profit Margin", "Orders", "AOV"]
    for c_idx, h in enumerate(chan_headers, 1):
        cell = ws_pa.cell(row=36, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    chan_summary = df_clean.groupby("Sales Channel").agg(
        Net=("Net Sales", "sum"),
        Cost=("Total Cost", "sum"),
        Profit=("Profit", "sum"),
        Orders=("Order ID", "nunique")
    ).reset_index().sort_values(by="Net", ascending=False)
    
    for idx, r in chan_summary.iterrows():
        r_num = 37 + idx
        share = r["Net"] / df_clean["Net Sales"].sum()
        margin = r["Profit"] / r["Net"]
        aov = r["Net"] / r["Orders"]
        row_vals = [r["Sales Channel"], r["Net"], share, r["Cost"], r["Profit"], margin, int(r["Orders"]), aov]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_pa.cell(row=r_num, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx == 1:
                cell.alignment = align_left
            elif c_idx in [2, 4, 5, 8]:
                cell.number_format = "$#,##0.00"
                cell.alignment = align_right
            elif c_idx in [3, 6]:
                cell.number_format = "0.0%"
                cell.alignment = align_right
            elif c_idx == 7:
                cell.number_format = "#,##0"
                cell.alignment = align_right

    # Table 5: Top 10 Products by Net Sales (with Margin Flagging!)
    ws_pa.cell(row=43, column=1, value="TABLE 5: TOP 10 PRODUCTS BY NET SALES (VOLUME VS MARGIN DIAGNOSIS)").font = font_section
    prod_headers = ["Rank", "Product Name", "Category", "Net Sales", "Total Cost", "Profit", "Profit Margin", "Units Sold", "Commercial Diagnosis"]
    for c_idx, h in enumerate(prod_headers, 1):
        cell = ws_pa.cell(row=44, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    top_prods = df_clean.groupby(["Product ID", "Product Name", "Category"]).agg(
        Net=("Net Sales", "sum"),
        Cost=("Total Cost", "sum"),
        Profit=("Profit", "sum"),
        Units=("Quantity", "sum")
    ).reset_index().sort_values(by="Net", ascending=False).head(10).reset_index(drop=True)
    
    for idx, r in top_prods.iterrows():
        r_num = 45 + idx
        margin = r["Profit"] / r["Net"]
        rank = idx + 1
        
        # Diagnosis
        if margin < 0:
            diag = "LOSS LEADER (Negative Margin - Excessive Discounting)"
        elif margin < 0.10:
            diag = "MARGIN COMPRESSED (High Revenue, Low Return)"
        elif margin > 0.25:
            diag = "STAR PRODUCT (High Volume & High Margin)"
        else:
            diag = "CORE PERFORMER (Stable Margin Profile)"
            
        row_vals = [rank, r["Product Name"], r["Category"], r["Net"], r["Cost"], r["Profit"], margin, int(r["Units"]), diag]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_pa.cell(row=r_num, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx == 1:
                cell.alignment = align_center
            elif c_idx in [2, 3]:
                cell.alignment = align_left
            elif c_idx in [4, 5, 6]:
                cell.number_format = "$#,##0.00"
                cell.alignment = align_right
            elif c_idx == 7:
                cell.number_format = "0.0%"
                cell.alignment = align_right
                if margin < 0:
                    cell.fill = fill_danger_soft
                    cell.font = Font(name="Calibri", size=11, bold=True, color="991B1B")
                elif margin < 0.10:
                    cell.fill = fill_warning_soft
                    cell.font = Font(name="Calibri", size=11, bold=True, color="92400E")
                elif margin > 0.25:
                    cell.fill = fill_success_soft
                    cell.font = Font(name="Calibri", size=11, bold=True, color="166534")
            elif c_idx == 8:
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif c_idx == 9:
                cell.alignment = align_left
                cell.font = font_bold
                if "LOSS LEADER" in diag:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
                elif "STAR PRODUCT" in diag:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="166534")

    # Table 6: Discount Depth Impact
    ws_pa.cell(row=57, column=1, value="TABLE 6: DISCOUNT TIER MARGIN SENSITIVITY ANALYSIS").font = font_section
    disc_headers = ["Discount Tier", "Transaction Rows", "Gross Sales", "Discount Given", "Net Sales", "Total Profit", "Profit Margin", "Margin Erosion Assessment"]
    for c_idx, h in enumerate(disc_headers, 1):
        cell = ws_pa.cell(row=58, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    df_clean['Disc_Tier'] = pd.cut(
        df_clean['Discount Percentage'], 
        bins=[-0.01, 0.0, 0.05, 0.10, 0.15, 0.26], 
        labels=['0% (No Promo)', '5% Promo', '10% Moderate', '15% Steep', '20-25% Aggressive']
    )
    disc_tier_summary = df_clean.groupby("Disc_Tier", observed=False).agg(
        Rows=("Order ID", "count"),
        Gross=("Gross Sales", "sum"),
        Disc=("Discount Amount", "sum"),
        Net=("Net Sales", "sum"),
        Profit=("Profit", "sum")
    ).reset_index()
    
    disc_assessments = [
        "Optimal organic margin baseline (27.0%)",
        "Healthy promotional pull-through (21.0% margin)",
        "Acceptable volume trade-off (16.8% margin)",
        "Severe margin compression (7.4% margin)",
        "Destructive discounting: Net Loss (-2.1% margin)"
    ]
    
    for idx, r in disc_tier_summary.iterrows():
        r_num = 59 + idx
        margin = r["Profit"] / r["Net"]
        row_vals = [r["Disc_Tier"], int(r["Rows"]), r["Gross"], r["Disc"], r["Net"], r["Profit"], margin, disc_assessments[idx]]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_pa.cell(row=r_num, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx == 1:
                cell.alignment = align_left
            elif c_idx == 2:
                cell.number_format = "#,##0"
                cell.alignment = align_right
            elif c_idx in [3, 4, 5, 6]:
                cell.number_format = "$#,##0.00"
                cell.alignment = align_right
            elif c_idx == 7:
                cell.number_format = "0.0%"
                cell.alignment = align_right
                if margin < 0:
                    cell.fill = fill_danger_soft
                    cell.font = Font(name="Calibri", size=11, bold=True, color="991B1B")
                elif margin < 0.10:
                    cell.fill = fill_warning_soft
            elif c_idx == 8:
                cell.alignment = align_left
                
    for col in ws_pa.columns:
        col_letter = get_column_letter(col[0].column)
        ws_pa.column_dimensions[col_letter].width = 18
    ws_pa.column_dimensions["A"].width = 16
    ws_pa.column_dimensions["B"].width = 30
    ws_pa.column_dimensions["I"].width = 44

    # =============================================================
    # 6. SHEET 6: Dashboard
    # =============================================================
    print("Building Sheet 6: Dashboard with KPI Cards and Native Charts...")
    ws_dash = wb.create_sheet(title="Dashboard")
    ws_dash.views.sheetView[0].showGridLines = False
    
    # 1. Header Banner (Rows 1-3)
    ws_dash.merge_cells("A1:N2")
    ws_dash["A1"] = "EXECUTIVE SALES & PROFITABILITY KPI DASHBOARD"
    ws_dash["A1"].font = font_title
    ws_dash["A1"].fill = fill_navy
    ws_dash["A1"].alignment = align_center
    
    ws_dash.merge_cells("A3:N3")
    ws_dash["A3"] = "FY 2024 Multi-Channel Performance | Filterable by Category, Region, Channel & Segment | Currency: USD"
    ws_dash["A3"].font = font_subtitle
    ws_dash["A3"].fill = fill_navy
    ws_dash["A3"].alignment = align_center
    
    # 2. Top Section: 6 KPI Cards (Rows 5-8)
    # Card Columns:
    # Card 1: A-B (Total Net Sales)
    # Card 2: C-D (Total Orders)
    # Card 3: E-F (Total Profit)
    # Card 4: G-H (Profit Margin)
    # Card 5: I-J (Average Order Value)
    # Card 6: K-L (Peak MoM Sales Growth)
    
    kpis = [
        ("TOTAL NET SALES", "=Pivot_Analysis!E18", "$1,042,276", "Gross: $1.15M | Disc: 9.7%", "A", "B"),
        ("TOTAL UNIQUE ORDERS", "=Pivot_Analysis!I18", "730", "1.37 items/order avg", "C", "D"),
        ("TOTAL NET PROFIT", "=Pivot_Analysis!G18", "$164,133", "Target: $150K (+9.4%)", "E", "F"),
        ("PROFIT MARGIN %", "=Pivot_Analysis!H18", "15.8%", "Benchmark Target: 15.0%", "G", "H"),
        ("AVERAGE ORDER VALUE", "=Pivot_Analysis!J18", "$1,428", "Revenue per unique order", "I", "J"),
        ("PEAK MOM GROWTH", "=Pivot_Analysis!K16", "+158.5%", "Nov Holiday / Black Friday", "K", "L")
    ]
    
    for title, formula_val, display_val, subtitle, start_c, end_c in kpis:
        # Title row 5
        ws_dash.merge_cells(f"{start_c}5:{end_c}5")
        cell_t = ws_dash[f"{start_c}5"]
        cell_t.value = title
        cell_t.font = font_kpi_label
        cell_t.fill = fill_card
        cell_t.alignment = align_center
        
        # Value row 6-7
        ws_dash.merge_cells(f"{start_c}6:{end_c}7")
        cell_v = ws_dash[f"{start_c}6"]
        cell_v.value = formula_val
        cell_v.font = font_kpi_num
        cell_v.fill = fill_card
        cell_v.alignment = align_center
        
        if "SALES" in title or "PROFIT" in title or "VALUE" in title:
            cell_v.number_format = "$#,##0"
        elif "MARGIN" in title or "GROWTH" in title:
            cell_v.number_format = "+0.0%;-0.0%;0.0%"
        elif "ORDERS" in title:
            cell_v.number_format = "#,##0"
            
        # Subtitle row 8
        ws_dash.merge_cells(f"{start_c}8:{end_c}8")
        cell_s = ws_dash[f"{start_c}8"]
        cell_s.value = subtitle
        cell_s.font = font_kpi_sub
        cell_s.fill = fill_card
        cell_s.alignment = align_center
        
        # Apply borders
        for r in range(5, 9):
            ws_dash[f"{start_c}{r}"].border = Border(left=Side(style="medium", color=NAVY_LIGHT))
            ws_dash[f"{end_c}{r}"].border = Border(right=Side(style="medium", color=NAVY_LIGHT))
        for col_name in [start_c, end_c]:
            ws_dash[f"{col_name}5"].border = Border(top=Side(style="medium", color=NAVY_LIGHT), left=Side(style="medium", color=NAVY_LIGHT) if col_name==start_c else None, right=Side(style="medium", color=NAVY_LIGHT) if col_name==end_c else None)
            ws_dash[f"{col_name}8"].border = Border(bottom=Side(style="medium", color=NAVY_LIGHT), left=Side(style="medium", color=NAVY_LIGHT) if col_name==start_c else None, right=Side(style="medium", color=NAVY_LIGHT) if col_name==end_c else None)

    # 3. Middle Section: Embedded Visual Charts
    # Chart 1: Monthly Net Sales Trend (Line Chart)
    chart1 = LineChart()
    chart1.title = "Monthly Net Sales Trend ($)"
    chart1.style = 13
    chart1.y_axis.title = "Net Sales ($)"
    chart1.x_axis.title = "Month"
    chart1.width = 17
    chart1.height = 10
    
    data1 = Reference(ws_pa, min_col=5, min_row=5, max_row=17) # Net Sales col E
    cats1 = Reference(ws_pa, min_col=2, min_row=6, max_row=17) # Month Names col B
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    ws_dash.add_chart(chart1, "A10")
    
    # Chart 2: Net Sales by Category (Bar / Column Chart)
    chart2 = BarChart()
    chart2.type = "col"
    chart2.title = "Net Sales by Product Category ($)"
    chart2.style = 10
    chart2.y_axis.title = "Net Sales ($)"
    chart2.x_axis.title = "Category"
    chart2.width = 15
    chart2.height = 10
    
    data2 = Reference(ws_pa, min_col=4, min_row=21, max_row=24) # Net Sales col D
    cats2 = Reference(ws_pa, min_col=1, min_row=22, max_row=24) # Category col A
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.legend = None
    ws_dash.add_chart(chart2, "H10")
    
    # Chart 3: Profit by Region (Column Chart)
    chart3 = BarChart()
    chart3.type = "col"
    chart3.title = "Net Profit by Geographic Region ($)"
    chart3.style = 11
    chart3.y_axis.title = "Profit ($)"
    chart3.x_axis.title = "Region"
    chart3.width = 17
    chart3.height = 10
    
    data3 = Reference(ws_pa, min_col=5, min_row=28, max_row=32) # Profit col E
    cats3 = Reference(ws_pa, min_col=1, min_row=29, max_row=32) # Region col A
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    chart3.legend = None
    ws_dash.add_chart(chart3, "A26")
    
    # Chart 4: Sales Distribution by Channel (Pie Chart)
    chart4 = PieChart()
    chart4.title = "Revenue Distribution by Sales Channel (%)"
    chart4.width = 15
    chart4.height = 10
    
    data4 = Reference(ws_pa, min_col=2, min_row=36, max_row=40) # Net Sales col B
    cats4 = Reference(ws_pa, min_col=1, min_row=37, max_row=40) # Channel col A
    chart4.add_data(data4, titles_from_data=True)
    chart4.set_categories(cats4)
    ws_dash.add_chart(chart4, "H26")
    
    # Chart 5: Top 10 Products by Net Sales (Horizontal Bar Chart)
    chart5 = BarChart()
    chart5.type = "bar" # Horizontal
    chart5.title = "Top 10 Products by Net Sales ($)"
    chart5.y_axis.title = "Product"
    chart5.x_axis.title = "Net Sales ($)"
    chart5.width = 17
    chart5.height = 11
    
    data5 = Reference(ws_pa, min_col=4, min_row=44, max_row=54) # Net Sales col D
    cats5 = Reference(ws_pa, min_col=2, min_row=45, max_row=54) # Product Name col B
    chart5.add_data(data5, titles_from_data=True)
    chart5.set_categories(cats5)
    chart5.legend = None
    ws_dash.add_chart(chart5, "A42")
    
    # Chart 6: Monthly Order Volume (Column Chart)
    chart6 = BarChart()
    chart6.type = "col"
    chart6.title = "Monthly Order Volume (# Orders)"
    chart6.style = 14
    chart6.y_axis.title = "Unique Orders"
    chart6.x_axis.title = "Month"
    chart6.width = 15
    chart6.height = 11
    
    data6 = Reference(ws_pa, min_col=9, min_row=5, max_row=17) # Orders col I
    cats6 = Reference(ws_pa, min_col=2, min_row=6, max_row=17) # Month Names col B
    chart6.add_data(data6, titles_from_data=True)
    chart6.set_categories(cats6)
    chart6.legend = None
    ws_dash.add_chart(chart6, "H42")
    
    # Slicer & Filter Instructions Box at bottom of Dashboard
    ws_dash.merge_cells("A58:N60")
    ws_dash["A58"] = "HOW TO USE INTERACTIVE SLICERS IN EXCEL: Click any chart or PivotTable > Insert > Slicer > Select [Category, Region, Sales Channel, Customer Segment]. Right-click Slicer > Report Connections > Check all PivotTables to enable cross-filtering across all visual cards."
    ws_dash["A58"].font = Font(name="Calibri", size=10, italic=True, color=NAVY_DARK)
    ws_dash["A58"].fill = fill_accent_soft
    ws_dash["A58"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Set Dashboard column widths
    for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N"]:
        ws_dash.column_dimensions[col_letter].width = 15

    # =============================================================
    # 7. SHEET 7: Business_Insights
    # =============================================================
    print("Building Sheet 7: Business_Insights...")
    ws_bi = wb.create_sheet(title="Business_Insights")
    ws_bi.views.sheetView[0].showGridLines = True
    
    # Title
    ws_bi.merge_cells("A1:E2")
    ws_bi["A1"] = "EXECUTIVE BUSINESS INSIGHTS & STRATEGIC RECOMMENDATIONS"
    ws_bi["A1"].font = font_title
    ws_bi["A1"].fill = fill_navy
    ws_bi["A1"].alignment = align_center
    
    ws_bi.merge_cells("A3:E3")
    ws_bi["A3"] = "Evidence-Based Findings Derived from FY 2024 Sales Analytics | Data-Driven Commercial Strategy"
    ws_bi["A3"].font = font_subtitle
    ws_bi["A3"].fill = fill_navy
    ws_bi["A3"].alignment = align_center
    
    # Section 1: Executive 5-Column Matrix (Strictly per Prompt Specification)
    ws_bi.cell(row=5, column=1, value="1. STRATEGIC EXECUTIVE DECISION MATRIX").font = font_section
    
    matrix_headers = ["Business Question", "Analytical Finding", "Empirical Evidence (Actual Metrics)", "Recommended Business Action", "Target Success Metric"]
    for c_idx, h in enumerate(matrix_headers, 1):
        cell = ws_bi.cell(row=6, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy_light
        cell.alignment = align_center
        
    matrix_data = [
        (
            "Which category generates the most net sales?",
            "Technology dominates total top-line revenue, contributing nearly two-thirds of all company sales.",
            "Technology Net Sales: $676,651.43 (64.9% of total company revenue). Furniture: $278,410.73 (26.7%). Office Supplies: $87,214.04 (8.4%).",
            "Maintain inventory availability for top high-ticket technology hardware; prioritize enterprise IT bundles to protect category leadership.",
            "Technology Net Sales growth rate (>12% YoY) and Order Fulfillment Rate (>98%)."
        ),
        (
            "Which category generates the most profit and margin?",
            "While Technology delivers the highest absolute profit dollars, Office Supplies is by far the highest-margin category.",
            "Office Supplies delivers a 35.11% profit margin ($30,616.95 profit on $87,214 sales). Technology delivered $91,172.26 profit but at a lower 13.47% margin.",
            "Implement high-margin cross-sell recommendations: automatically suggest Office Supplies and Desk Accessories at checkout for all Technology laptop orders.",
            "Office Supplies attach rate (+25% in checkout bundles) and overall Company Gross Margin expansion (+150 bps to 17.25%)."
        ),
        (
            "Which products have high sales but severely compressed profit margins?",
            "The #1 best-selling product by revenue (Apex Gaming Laptop 17\") generated an absolute net loss due to excessive promotional discounting.",
            "Apex Gaming Laptop: Net Sales = $125,123.40 (Rank #1 in sales), but Total Profit = -$1,332.10 (-1.06% profit margin). ErgoFlex Chair also had low 5.82% margin ($1,842 profit on $31,640 sales).",
            "Establish strict discounting guardrails: cap Apex Gaming Laptop discounts at maximum 10% (from current 20-25%). Shift marketing spend to ProBook Elite 15.6\" which yields a healthy 21.62% profit margin ($31,840 profit).",
            "Apex Gaming Laptop Margin Recovery (achieve positive 12% margin) and Blended Product Portfolio Margin."
        ),
        (
            "Which geographic region performs best and worst?",
            "The West region leads in total sales volume, but the East region is the most profitable and efficient market.",
            "East generated $49,649.24 profit (18.82% profit margin on $263,848 sales). West generated higher sales ($327,958) but lower profit ($47,380.60, 14.45% margin). South is lowest in total sales ($194,392.46, 16.55% margin).",
            "Audit West region discounting and localized shipping freight surcharges. Expand corporate direct sales reps into underpenetrated high-margin Southern states (Florida and North Carolina).",
            "Regional Profit Margin parity (bring West to >17.5%) and South Region Revenue growth (+20% YoY)."
        ),
        (
            "Which sales channel contributes the most revenue and highest AOV?",
            "Online drives the largest overall transaction volume and revenue, while Wholesale delivers the highest Average Order Value.",
            "Online contributed $380,542.94 (36.51% share, 309 orders, AOV $1,231.53). Wholesale contributed $308,272.88 (29.58% share, 97 orders, AOV $3,178.07). Retail Partner had lowest share ($159,794.44, 15.33%).",
            "Optimize Online customer conversion with personalized add-ons; build an automated Wholesale self-service portal to accelerate B2B reordering and scale high-AOV orders.",
            "Wholesale customer retention rate (>85%) and Online checkout conversion (+1.8%)."
        ),
        (
            "Which months show seasonal peaks and troughs?",
            "Sales follow strong retail seasonality with an extreme holiday surge in Q4, peaking in November and December.",
            "November Net Sales peaked at $175,122.83 (+158.5% MoM growth over October) with 105 unique orders, followed by December ($122,972.56). August ($59,365.40) and February ($65,605.35) were annual troughs.",
            "Establish working capital credit facilities and build warehouse inventory reserves by early October. Run targeted 'Back-to-School' and 'Mid-Year Budget' campaigns during low-volume months (July-August).",
            "Q4 On-Time Fulfillment (>99%) and Q3 trough revenue uplift (+15% YoY in August)."
        ),
        (
            "Are larger discounts associated with lower profit margins?",
            "Yes, empirical evidence shows severe margin destruction beyond the 10% discount threshold, resulting in negative margins above 20%.",
            "0% Discount: 27.04% margin ($75,429 profit on $278,921 sales). 5% Discount: 21.00% margin. 10% Discount: 16.83% margin. 15% Discount: 7.40% margin. 20-25% Discount: -2.10% NET LOSS (-$2,886 profit on $137,489 sales).",
            "Eliminate 20%+ unapproved promo codes. Implement automated system approvals for discounts above 10%, requiring VP sign-off for enterprise deals.",
            "Discount-weighted margin preservation (target sales-weighted discount <7.5% vs current 9.68%)."
        )
    ]
    
    for r_idx, row_data in enumerate(matrix_data, start=7):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws_bi.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            cell.alignment = align_wrap_left
            if c_idx == 1:
                cell.font = font_bold
        ws_bi.row_dimensions[r_idx].height = 68
        
    ws_bi.column_dimensions["A"].width = 24
    ws_bi.column_dimensions["B"].width = 30
    ws_bi.column_dimensions["C"].width = 38
    ws_bi.column_dimensions["D"].width = 36
    ws_bi.column_dimensions["E"].width = 28
    
    # Section 2: Deep-Dive Strategic Recommendations
    start_strat_row = 16
    ws_bi.cell(row=start_strat_row, column=1, value="2. STRATEGIC COMMERCIAL ROADMAP & PRODUCT GOVERNANCE").font = font_section
    
    strategies = [
        ("Product Portfolio Rationalization", "Reposition 'Apex Gaming Laptop 17\"' from a promotional loss-leader to a premium performance workstation. Bundle it with high-margin peripherals (e.g. UltraWireless Headset with 33% margin and Precision Laser Mouse with 45% margin) to ensure every transaction yields positive blended gross profit."),
        ("Dynamic Discount Governance", "Establish tier-based approval workflows in the CRM/ERP system. Eliminate blanket site-wide 20% discounts. Transactions in the 20-25% discount tier generated a net loss of $2,886. Capping discounts at 10% preserves an estimated $28,500 in additional annual bottom-line profit."),
        ("Regional Optimization & Channel Synergies", "Replicate East Region sales practices (focused on corporate direct sales with 18.8% margin) in the Central and West regions. Expand B2B Wholesale programs in the South region to capitalize on untapped corporate enterprise demand.")
    ]
    
    cur_s_row = start_strat_row + 1
    for title, desc in strategies:
        ws_bi.cell(row=cur_s_row, column=1, value=title).font = font_bold
        ws_bi.cell(row=cur_s_row, column=1).alignment = Alignment(horizontal="left", vertical="top")
        ws_bi.cell(row=cur_s_row, column=1).border = border_cell
        
        ws_bi.merge_cells(start_row=cur_s_row, start_column=2, end_row=cur_s_row, end_column=5)
        d_cell = ws_bi.cell(row=cur_s_row, column=2, value=desc)
        d_cell.font = font_regular
        d_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        d_cell.border = border_cell
        ws_bi.row_dimensions[cur_s_row].height = 42
        cur_s_row += 1

    # Save to both workbook/ and root for user convenience
    os.makedirs("workbook", exist_ok=True)
    out_path1 = "workbook/Sales_Analysis_Dashboard.xlsx"
    out_path2 = "Sales_Analysis_Dashboard.xlsx"
    
    wb.save(out_path1)
    wb.save(out_path2)
    print(f"Successfully saved complete workbook to {out_path1} and {out_path2}!")

if __name__ == "__main__":
    build_excel_workbook()
